"""
Comprehensive Verification Test Suite for Marksheet Digitizer
Continuous Smartphone Batch-Scanning Cloud System & Subject Router.
Tests:
  1. API Status & Subject Summary Endpoints
  2. Multi-photo Smartphone Batch Upload Simulation
  3. Subject Lock Override (Force 24MCPC517 / 24MCPC521)
  4. Deduplication & Upsert Logic
  5. Standalone Subject Excel Download & USN Sorting
  6. Master Consolidated Excel Workbook (4 sheets, 0 duplicates, 100% Z3 arithmetic validity)
"""

import io
import sys
import logging
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import pandas as pd
import openpyxl
from fastapi.testclient import TestClient

from server import app
from src.config import RAW_DATA_DIR, VALID_COURSE_CODES

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("test_suite")


def run_tests():
    print("=" * 80)
    print("RUNNING COMPREHENSIVE BATCH SCANNING & SUBJECT ROUTING TEST SUITE")
    print("=" * 80)

    client = TestClient(app)

    # -------------------------------------------------------------------------
    # TEST 1: System Status
    # -------------------------------------------------------------------------
    print("\n[TEST 1] Testing GET /api/status...")
    resp = client.get("/api/status")
    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
    status_data = resp.json()
    assert status_data["status"] == "online"
    assert "device" in status_data
    assert "subjects" in status_data
    for code in ["24MCPC522", "24MCPC517", "24MCPC521"]:
        assert code in status_data["subjects"], f"Missing {code} in status subjects"
    print(f"  [PASS] System Status OK: Device={status_data['device']}, Records={status_data['total_records']}")

    # -------------------------------------------------------------------------
    # TEST 2: Subject Summary Endpoint
    # -------------------------------------------------------------------------
    print("\n[TEST 2] Testing GET /api/subjects/summary...")
    resp = client.get("/api/subjects/summary")
    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}"
    summary = resp.json()
    assert summary["success"] is True
    assert "subjects" in summary
    for code in ["24MCPC522", "24MCPC517", "24MCPC521"]:
        assert code in summary["subjects"]
        subj = summary["subjects"][code]
        assert "student_count" in subj
        assert "average_cie" in subj
        assert "records" in subj
        print(f"  [PASS] Subject {code}: {subj['student_count']} students | Avg CIE: {subj['average_cie']}/50")

    # -------------------------------------------------------------------------
    # TEST 3: Multi-Photo Smartphone Batch Upload Simulation (Auto-Detect)
    # -------------------------------------------------------------------------
    print("\n[TEST 3] Testing POST /api/booklets/batch-upload (Multi-photo upload)...")
    sample_files = sorted(RAW_DATA_DIR.glob("booklet_*_p1.jpg"))[:2]
    assert len(sample_files) >= 2, "Need at least 2 sample files in data/raw/"

    files_payload = []
    for f in sample_files:
        files_payload.append(
            ("files", (f.name, f.read_bytes(), "image/jpeg"))
        )

    resp = client.post(
        "/api/booklets/batch-upload",
        files=files_payload,
        data={"forced_course_code": "", "subject_folder": "live_mobile_test"},
    )
    assert resp.status_code == 200, f"Batch upload failed: {resp.text}"
    batch_data = resp.json()
    assert batch_data["success"] is True
    assert batch_data["count"] == 2
    assert len(batch_data["results"]) == 2

    for r in batch_data["results"]:
        assert r["usn"].startswith("4MT24MC"), f"Invalid USN prefix: {r['usn']}"
        assert r["course_code"] in VALID_COURSE_CODES, f"Invalid course code: {r['course_code']}"
        assert r["total_cie_50"] >= 0 and r["total_cie_50"] <= 50
        assert r["image_url"].startswith("/cloud-vault/")
        print(f"  [PASS] Uploaded & Routed: [{r['course_code']}] USN={r['usn']} | CIE={r['total_cie_50']}/50 | Status={r['status']}")

    # -------------------------------------------------------------------------
    # TEST 4: Forced Subject Lock Override
    # -------------------------------------------------------------------------
    print("\n[TEST 4] Testing Subject Lock Override (Force 24MCPC517)...")
    f_test = sample_files[0]
    resp = client.post(
        "/api/booklets/batch-upload",
        files=[("files", (f_test.name, f_test.read_bytes(), "image/jpeg"))],
        data={"forced_course_code": "24MCPC517", "subject_folder": "forced_test"},
    )
    assert resp.status_code == 200, f"Forced upload failed: {resp.text}"
    forced_res = resp.json()["results"][0]
    assert forced_res["course_code"] == "24MCPC517", f"Expected 24MCPC517, got {forced_res['course_code']}"
    print(f"  [PASS] Forced Subject Lock OK: Routed to {forced_res['course_code']} (USN: {forced_res['usn']})")

    # -------------------------------------------------------------------------
    # TEST 5: Deduplication & Upsert Verification
    # -------------------------------------------------------------------------
    print("\n[TEST 5] Testing Deduplication & Upsert on Duplicate USN Upload...")
    sum_before = client.get("/api/subjects/summary").json()["subjects"]["24MCPC517"]["student_count"]

    # Re-upload the exact same booklet with forced 24MCPC517
    resp_dup = client.post(
        "/api/booklets/batch-upload",
        files=[("files", (f_test.name, f_test.read_bytes(), "image/jpeg"))],
        data={"forced_course_code": "24MCPC517", "subject_folder": "forced_test_dup"},
    )
    assert resp_dup.status_code == 200

    sum_after = client.get("/api/subjects/summary").json()["subjects"]["24MCPC517"]["student_count"]
    assert sum_after == sum_before, f"Duplicate created! Count before={sum_before}, after={sum_after}"
    print(f"  [PASS] Deduplication Verified: Student count remained {sum_after} without duplicate row.")

    # -------------------------------------------------------------------------
    # TEST 6: Standalone Subject Excel Download & USN Sorting
    # -------------------------------------------------------------------------
    print("\n[TEST 6] Testing Standalone Subject Excel Download (GET /api/subjects/24MCPC522/download)...")
    resp_subj_excel = client.get("/api/subjects/24MCPC522/download")
    assert resp_subj_excel.status_code == 200
    assert "application/vnd.openxmlformats" in resp_subj_excel.headers["content-type"]

    wb_subj = openpyxl.load_workbook(io.BytesIO(resp_subj_excel.content))
    ws_subj = wb_subj.active
    rows_subj = list(ws_subj.iter_rows(values_only=True))
    assert len(rows_subj) >= 1, "Worksheet is empty"
    header_subj = rows_subj[0]
    assert "Student USN" in header_subj
    assert "Total CIE (50)" in header_subj

    data_rows = rows_subj[1:]
    usns = [r[0] for r in data_rows if r[0]]
    sorted_usns = sorted(usns)
    assert usns == sorted_usns, f"USNs are not in ascending order! Got {usns}"
    print(f"  [PASS] Standalone Excel OK: {len(usns)} rows cleanly formatted and strictly sorted by USN.")

    # -------------------------------------------------------------------------
    # TEST 7: Master Consolidated Multi-Sheet Workbook & Arithmetic Constraints
    # -------------------------------------------------------------------------
    print("\n[TEST 7] Testing Consolidated Master Workbook (GET /api/booklets/download-master)...")
    resp_master = client.get("/api/booklets/download-master")
    assert resp_master.status_code == 200
    assert "application/vnd.openxmlformats" in resp_master.headers["content-type"]

    wb_master = openpyxl.load_workbook(io.BytesIO(resp_master.content))
    sheet_names = wb_master.sheetnames
    expected_sheets = ["All_Records_Combined", "24MCPC522", "24MCPC517", "24MCPC521"]
    for s in expected_sheets:
        assert s in sheet_names, f"Missing sheet '{s}' in master workbook. Found: {sheet_names}"
    print(f"  [PASS] Master Workbook Sheets Verified: {sheet_names}")

    # Verify arithmetic constraints across all records in all sheets
    total_checked = 0
    for s_name in sheet_names:
        ws = wb_master[s_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) <= 1:
            continue
        hdr = list(rows[0])
        col_map = {col_name: idx for idx, col_name in enumerate(hdr)}

        sheet_usns = []
        for r in rows[1:]:
            usn = r[col_map["Student USN"]]
            if not usn:
                continue
            sheet_usns.append(usn)

            t1 = int(r[col_map["Test 1 (15)"]] or 0)
            t2 = int(r[col_map["Test 2 (15)"]] or 0)
            aba = int(r[col_map["ABA (10)"]] or 0)
            th = int(r[col_map["Theory (40)"]] or 0)
            viva = int(r[col_map["Viva (5)"]] or 0)
            lab = int(r[col_map["Lab (5)"]] or 0)
            pr = int(r[col_map["Pract (10)"]] or 0)
            cie = int(r[col_map["Total CIE (50)"]] or 0)

            # Invariants
            assert t1 + t2 + aba == th, f"Constraint C1 violated for {usn}: {t1}+{t2}+{aba} != {th}"
            assert viva + lab == pr, f"Constraint C2 violated for {usn}: {viva}+{lab} != {pr}"
            assert th + pr == cie, f"Constraint C3 violated for {usn}: {th}+{pr} != {cie}"
            total_checked += 1

        # Check deduplication per subject sheet
        if s_name != "All_Records_Combined":
            assert len(sheet_usns) == len(set(sheet_usns)), f"Duplicate USNs in sheet {s_name}: {sheet_usns}"
            assert sheet_usns == sorted(sheet_usns), f"Sheet {s_name} not sorted by USN!"

    print(f"  [PASS] 100% Z3 Arithmetic Validity Verified across all {total_checked} rows in workbook!")
    print(f"  [PASS] Zero duplicate USNs per subject sheet.")
    print(f"  [PASS] Ascending USN sorting verified.")

    print("\n" + "=" * 80)
    print("ALL 7 TEST SUITES PASSED WITH 100% SUCCESS!")
    print("=" * 80)


if __name__ == "__main__":
    run_tests()
